#!/usr/bin/env python3
"""Prepare publication-facing S1–S7 tables and one styled workbook.
import os
Source data are frozen v10 audit outputs. Move legacy comparison to internal audit.
"""
from pathlib import Path
import csv,shutil,json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

ROOT=Path(os.environ.get("LOQUAT_PROJECT_ROOT",str(Path.home()/"Desktop/LOQUAT_DWARFING_PUBLIC"))).expanduser().resolve()
A=ROOT/"11_methods_audit_20261020"
O=ROOT/"12_submission_v10_assets"
def read(p):
    with p.open(newline="",encoding="utf-8-sig") as f:return list(csv.DictReader(f))
def save(p,rows,fields):
    with p.open("w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

old=O/"Table_S7_Legacy_vs_V10.csv"
if old.exists():shutil.copy2(old,O/"AUDIT_Legacy_vs_V10_not_for_submission.csv")
r=read(A/"KEGG_v10_flavonoid_summary.csv")
keys={"s0.7_b0_pool":(".70","collinearity passed"),"s0.7_b0_external":(".70","external support"),
"s0.7_b0_both":(".70","both external tissues"),
"s0.8_b0_pool":(".80","collinearity passed"),"s0.8_b0_external":(".80","external support"),"s0.8_b0_both":(".80","both external tissues"),
"s0.9_b0_pool":(".90","collinearity passed"),"s0.9_b0_external":(".90","external support"),"s0.9_b0_both":(".90","both external tissues")}
v=[]
for x in r:
    if x["id"] not in keys:continue
    t,kind=keys[x["id"]]
    v.append({"Collinearity_threshold":float(t),"Candidate_subset":kind,
    "Candidate_n":int(x["original_n"]),"KEGG_mapped_candidate_n":int(x["mapped_n"]),
    "Flavonoid_hits":int(x["hits"]),"Flavonoid_fold_enrichment":float(x["fold"]),
    "Flavonoid_BH_FDR_151pathways":float(x["FDR_BH_151"])})
v.sort(key=lambda x:(x["Collinearity_threshold"],["collinearity passed","external support","both external tissues"].index(x["Candidate_subset"])))
save(O/"Table_S7_Collinearity_threshold_sensitivity.csv",v,list(v[0]))
if old.exists():old.unlink()

files=[
("S1_Eligible269",O/"Table_S1_All_269_eligible.csv",269),
("S2_Collinearity190",O/"Table_S2_Synteny_190.csv",190),
("S3_External24",O/"Table_S3_External_24.csv",24),
("S4_BothTissues5",O/"Table_S4_BothTissues_5.csv",5),
("S5_KEGG151",O/"Table_S5_KEGG_fixed151_all_sets.csv",453),
("S6_PKS_Sensitivity",O/"Table_S6_PKS_paralog_sensitivity.csv",12),
("S7_Threshold",O/"Table_S7_Collinearity_threshold_sensitivity.csv",9),
]
for _,p,n in files:
    assert len(read(p))==n,(p,n,len(read(p)))
wb=Workbook();ws=wb.active;ws.title="READ_ME"
intro=[
("LOQUAT | Supplementary Tables (v10)", "Provisional research tables — consistent with the v10 candidate-selection and KEGG audit"),
("Dataset", "Apple/pear/loquat strict RBH correspondence; public expression contrasts; loquat TBY stem TPM"),
("Candidate selection", "20,795 strict RBH → 269 expression-supported → 190 local-collinearity → 24 external → 5 both-tissue"),
("Collinearity", "Minimum(pear, loquat) orientation-free anchor-order fraction >= 0.80; at least 8 apple reference anchors"),
("KEGG", "One-sided hypergeometric ORA; 4,775 annotated strict-RBH genes; fixed 151 pathway BH correction; frozen 2026-10-20 mappings"),
("Uncertainty", "Primary apple RNA-extraction replicates need author clarification about independent plant sampling pools"),
("Terminology", "The 5-gene category indicates significant same-sign effects in both external tissues, not cross-experiment sign replication."),
("Document status", "Internal pre-submission analytical supplement; do not submit before manuscript-wide editorial and dataset audit."),
("Source scripts", "12_build_supp_v10.py; 12_build_submission_tables_v10.py; v10_universal_synteny.py; KEGG_v10_threshold_sensitivity.R"),
("Sources", "Nature Genetics 2024, DOI 10.1038/s41588-024-01657-2, BioProject PRJNA814760; PRJNA1483470; PRJCA059134; PRJNA1166645"),
]
for i,(k,value) in enumerate(intro,1):
    ws.cell(i,1,k)
    ws.cell(i,2,value)
    ws.cell(i,1).font=Font(bold=True,color="24464C",size=11)
    ws.cell(i,2).alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[i].height=35 if i in (1,3,4,5,6,7,8,9,10) else 25
ws.column_dimensions["A"].width=27
ws.column_dimensions["B"].width=116
ws.freeze_panes="A2"
for name,p,n in files:
    records=read(p)
    sheet=wb.create_sheet(name)
    hdr=list(records[0]);sheet.append(hdr)
    for row in records:
        vals=[]
        for key in hdr:
            val=row.get(key,"")
            if val in ("",None):vals.append(None);continue
            if val in ("True","False"):vals.append(val=="True");continue
            if key.endswith("ID") or key.endswith("GeneID") or key in ("Apple_GeneID","Pear_GeneID","Loquat_GeneID","KEGG_ID","Candidate_Set","Candidate_subset","PKS5_Exclusion","Stage_Response_Class","Apple_Annotation_postselection"):
                vals.append(val);continue
            try:
                vals.append(float(val) if any(c in val for c in ".eE") else int(val))
            except (ValueError,TypeError):vals.append(val)
        sheet.append(vals)
    sheet.freeze_panes="D2" if len(hdr)>6 else "A2"
    sheet.auto_filter.ref=sheet.dimensions
    sheet.sheet_view.showGridLines=False
    for c in sheet[1]:
        c.font=Font(name="Aptos",bold=True,color="FFFFFF",size=10)
        c.fill=PatternFill(fill_type="solid",fgColor="215A59")
        c.alignment=Alignment(wrap_text=True,vertical="center")
    sheet.row_dimensions[1].height=45
    maxrows=sheet.max_row
    for col_ix,key in enumerate(hdr,1):
        letter=get_column_letter(col_ix)
        width=38 if "Annotation" in key or "Category" in key else 21 if ("Gene" in key or "FDR" in key) else 18
        if key in ["Candidate_Set","Candidate_subset","PKS5_Exclusion","Stage_Response_Class"]:width=26
        if key in ["Apple_GeneID","Pear_GeneID","Loquat_GeneID"]:width=20
        sheet.column_dimensions[letter].width=min(width,41)
    for ri in range(2,maxrows+1):
        if ri%2==0:
            for cell in sheet[ri]:
                cell.fill=PatternFill(fill_type="solid",fgColor="F5F9F9")
        for cell in sheet[ri]:
            cell.font=Font(name="Aptos",size=9,color="23363D")
            if isinstance(cell.value,float):cell.number_format="0.0000" if abs(cell.value)>=.0001 else "0.00E+00"
wb.active=0
out=O/"Loquat_Supplementary_Tables_v10_PROVISIONAL.xlsx"
wb.save(out)
chk=load_workbook(out,read_only=True,data_only=True)
for name,p,n in files:assert chk[name].max_row-1==n
assert len(chk.sheetnames)==8
print("created:",out, "sheet_counts:", {name:len(read(p)) for name,p,n in files})
